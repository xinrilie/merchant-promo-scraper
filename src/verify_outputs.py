"""Validate the expected artifacts from a completed Merchant Promo Scraper run.

The project root is resolved from this file's location, not interpolated into a
Python command string. This avoids Windows quoting failures for paths containing
spaces or a trailing backslash.
"""
from __future__ import annotations

import argparse
import csv
import json
import sys
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parent.parent

REQUIRED_FILES = {
    "workbook": Path("output/Katalog_Promo_Final.xlsx"),
    "catalog_csv": Path("output/Katalog_Promo_Final.csv"),
    "image_index": Path("runtime/promo_multisource_catalog_images.csv"),
    "source_health": Path("runtime/source_health_summary.json"),
}


def _check_csv(path: Path) -> str | None:
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            row = next(csv.reader(handle), None)
        if not row or not any(str(cell).strip() for cell in row):
            return "CSV kosong atau tidak memiliki header."
    except (OSError, UnicodeError, csv.Error) as exc:
        return f"CSV tidak dapat dibaca: {exc}"
    return None


def verify_outputs(root: Path = ROOT) -> dict[str, Any]:
    """Return a report with checked files, workbook sheets, and errors."""
    root = Path(root).resolve()
    report: dict[str, Any] = {"root": str(root), "files": {}, "sheets": [], "errors": []}
    errors: list[str] = report["errors"]
    resolved: dict[str, Path] = {}

    for key, relative_path in REQUIRED_FILES.items():
        path = root / relative_path
        resolved[key] = path
        if not path.is_file():
            errors.append(f"File {key} tidak ditemukan: {path}")
            report["files"][key] = {"path": str(path), "ok": False}
            continue
        size = path.stat().st_size
        if size <= 0:
            errors.append(f"File {key} kosong: {path}")
            report["files"][key] = {"path": str(path), "size_bytes": size, "ok": False}
            continue
        report["files"][key] = {"path": str(path), "size_bytes": size, "ok": True}

    workbook_path = resolved["workbook"]
    if report["files"]["workbook"]["ok"]:
        try:
            from openpyxl import load_workbook

            workbook = load_workbook(workbook_path, read_only=True, data_only=True)
            report["sheets"] = list(workbook.sheetnames)
            workbook.close()
            if not report["sheets"]:
                errors.append(f"Workbook tidak memiliki sheet: {workbook_path}")
        except Exception as exc:  # malformed/corrupt workbook should be reported cleanly
            errors.append(f"Workbook tidak dapat dibuka: {workbook_path} ({exc})")
            report["files"]["workbook"]["ok"] = False

    for key in ("catalog_csv", "image_index"):
        if report["files"][key]["ok"]:
            problem = _check_csv(resolved[key])
            if problem:
                errors.append(f"{problem} File: {resolved[key]}")
                report["files"][key]["ok"] = False

    health_path = resolved["source_health"]
    if report["files"]["source_health"]["ok"]:
        try:
            payload = json.loads(health_path.read_text(encoding="utf-8-sig"))
            if not isinstance(payload, dict):
                raise ValueError("root JSON harus berupa objek")
        except (OSError, UnicodeError, json.JSONDecodeError, ValueError) as exc:
            errors.append(f"Laporan source health tidak valid: {health_path} ({exc})")
            report["files"]["source_health"]["ok"] = False

    report["ok"] = not errors
    return report


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--root",
        type=Path,
        default=ROOT,
        help="Root folder proyek (default: folder induk src/)",
    )
    args = parser.parse_args(argv)
    report = verify_outputs(args.root)

    print(f"Root proyek: {report['root']}")
    for key, info in report["files"].items():
        status = "OK" if info.get("ok") else "GAGAL"
        size = info.get("size_bytes")
        suffix = f" ({size} bytes)" if size is not None else ""
        print(f"{status}: {key}: {info['path']}{suffix}")
    if report["sheets"]:
        print("Sheet: " + ", ".join(report["sheets"]))

    if report["errors"]:
        for error in report["errors"]:
            print(f"ERROR: {error}", file=sys.stderr)
        return 1

    print("Verifikasi output berhasil.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
